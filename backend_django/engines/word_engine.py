import os
import time
import uuid
import zipfile
import logging
import re
import io
import tempfile
import shutil
from pathlib import Path
from typing import Optional, Dict, Any, Tuple

import pymupdf
from pdf2docx import Converter
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fontTools.ttLib import TTFont
import xml.etree.ElementTree as ET
from django.conf import settings

logger = logging.getLogger(__name__)


class WordEngineException(Exception):
    """Base exception for Word engine failures."""
    pass


class InvalidPdfError(WordEngineException):
    """Raised when the uploaded file is not a valid PDF."""
    pass


class PasswordProtectedPdfError(WordEngineException):
    """Raised when the PDF requires a password to open."""
    pass


def parse_cmap_clean(text: str) -> Dict[int, str]:
    """
    Parses an Adobe-Identity-UCS or ToUnicode CMap stream into a dictionary
    mapping integer Character IDs (CIDs) to Unicode character strings.
    Handles bfchar, bfrange with list destinations, and bfrange with sequential offsets.
    """
    mapping: Dict[int, str] = {}
    in_bfchar = False
    in_bfrange = False

    for line in text.splitlines():
        line = line.strip()
        if not line or line.startswith('%'):
            continue
        if 'beginbfchar' in line:
            in_bfchar = True
            in_bfrange = False
            continue
        elif 'endbfchar' in line:
            in_bfchar = False
            continue
        elif 'beginbfrange' in line:
            in_bfrange = True
            in_bfchar = False
            continue
        elif 'endbfrange' in line:
            in_bfrange = False
            continue

        if in_bfchar:
            tokens = [t.strip('<>') for t in line.split() if t.startswith('<') and t.endswith('>')]
            if len(tokens) == 2:
                try:
                    src = int(tokens[0], 16)
                    dst_hex = tokens[1]
                    chars = ''.join(chr(int(dst_hex[i:i+4], 16)) for i in range(0, len(dst_hex), 4))
                    mapping[src] = chars
                except Exception:
                    pass
        elif in_bfrange:
            if '[' in line and ']' in line:
                pre = line[:line.index('[')].strip()
                tokens = [t.strip('<>') for t in pre.split() if t.startswith('<') and t.endswith('>')]
                bracket_content = line[line.index('[')+1:line.index(']')].strip()
                items = [t.strip('<>') for t in bracket_content.split() if t.startswith('<') and t.endswith('>')]
                if len(tokens) == 2 and len(items) > 0:
                    try:
                        start = int(tokens[0], 16)
                        for i, item in enumerate(items):
                            chars = ''.join(chr(int(item[j:j+4], 16)) for j in range(0, len(item), 4))
                            mapping[start + i] = chars
                    except Exception:
                        pass
            else:
                tokens = [t.strip('<>') for t in line.split() if t.startswith('<') and t.endswith('>')]
                if len(tokens) == 3:
                    try:
                        start = int(tokens[0], 16)
                        end = int(tokens[1], 16)
                        dst_hex = tokens[2]
                        if len(dst_hex) == 4:
                            base = int(dst_hex, 16)
                            for code in range(start, end + 1):
                                mapping[code] = chr(base + (code - start))
                        elif len(dst_hex) == 8:
                            u1 = int(dst_hex[:4], 16)
                            u2 = int(dst_hex[4:], 16)
                            for code in range(start, end + 1):
                                mapping[code] = chr(u1) + chr(u2 + (code - start))
                    except Exception:
                        pass
    return mapping


def repair_pdf_cmaps(pdf_path: str) -> Tuple[str, bool]:
    """
    Inspects embedded Type0 fonts in the PDF. When fonts use Identity-H encoding
    (where CID == GID) but have corrupted or incomplete ToUnicode CMaps (common in
    Microsoft Word PDF exports for Indic scripts like Tamil), reconstructs an accurate
    ToUnicode CMap from the embedded TrueType glyph outline names ('uni0Bxx').
    Returns (path_to_use, was_modified).
    """
    try:
        doc = pymupdf.open(pdf_path)
    except Exception as err:
        logger.warning(f"Could not open PDF for font inspection: {err}")
        return pdf_path, False

    modified = False
    font_xrefs = set()
    for page in doc:
        for f in page.get_fonts():
            font_xrefs.add(f[0])

    for f_xref in font_xrefs:
        try:
            font_obj = doc.xref_object(f_xref)
            if '/Type0' not in font_obj or '/ToUnicode' not in font_obj:
                continue

            to_uni_match = re.search(r'/ToUnicode\s+(\d+)\s+0\s+R', font_obj)
            if not to_uni_match:
                continue
            to_uni_xref = int(to_uni_match.group(1))

            font_info = doc.extract_font(f_xref)
            ext, data = font_info[1], font_info[3]
            if ext not in ['ttf', 'otf'] or not data:
                continue

            tt = TTFont(io.BytesIO(data))
            gorder = tt.getGlyphOrder()

            # Check if font contains Indic/Tamil glyphs (e.g. uni0Bxx)
            has_indic_glyphs = any(g.startswith('uni0B') and len(g) == 7 for g in gorder)
            if not has_indic_glyphs:
                continue

            old_cmap_stream = doc.xref_stream(to_uni_xref).decode('latin1', errors='ignore')
            existing_map = parse_cmap_clean(old_cmap_stream)

            final_map: Dict[int, str] = {}
            for gid in range(len(gorder)):
                gname = gorder[gid]
                if gname.startswith('uni0B') and len(gname) == 7:
                    try:
                        final_map[gid] = chr(int(gname[3:], 16))
                    except Exception:
                        pass
                elif gname == 'space':
                    final_map[gid] = ' '
                elif gname in ['hyphen', 'endash']:
                    final_map[gid] = '-'
                elif gid in existing_map:
                    # Preserve ligatures/conjuncts from original CMap
                    final_map[gid] = existing_map[gid]

            diff_count = sum(1 for g, val in final_map.items() if existing_map.get(g) != val)
            if diff_count > 0:
                bfchars = []
                for gid, val in sorted(final_map.items()):
                    val_hex = ''.join(f'{ord(c):04X}' for c in val)
                    bfchars.append(f'<{gid:04X}> <{val_hex}>')

                cmap_lines = [
                    '/CIDInit /ProcSet findresource begin',
                    '12 dict begin',
                    'begincmap',
                    '/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def',
                    '/CMapName /Adobe-Identity-UCS def',
                    '/CMapType 2 def',
                    '1 begincodespacerange',
                    '<0000> <FFFF>',
                    'endcodespacerange'
                ]
                for i in range(0, len(bfchars), 100):
                    chunk = bfchars[i:i+100]
                    cmap_lines.append(f'{len(chunk)} beginbfchar')
                    cmap_lines.extend(chunk)
                    cmap_lines.append('endbfchar')
                cmap_lines.extend([
                    'endcmap',
                    'CMapName currentdict /CMap defineresource pop',
                    'end',
                    'end'
                ])
                new_cmap_str = '\n'.join(cmap_lines)
                doc.update_stream(to_uni_xref, new_cmap_str.encode('latin1'))
                modified = True
                logger.info(f"Repaired font xref {f_xref} (CMap {to_uni_xref}): {diff_count} glyph mappings corrected.")
        except Exception as exc:
            logger.warning(f"Notice while inspecting font xref {f_xref}: {exc}")

    if modified:
        tmp_fd, tmp_path = tempfile.mkstemp(suffix='.pdf')
        os.close(tmp_fd)
        doc.save(tmp_path, garbage=3, deflate=True)
        doc.close()
        return tmp_path, True

    doc.close()
    return pdf_path, False


def normalize_indic_text(text: str) -> str:
    """
    Normalizes Indic script text extracted in visual order into standard logical Unicode order.
    Specifically handles Tamil pre-base vowels (U+0BC6, U+0BC7, U+0BC8) and combining signs.
    """
    if not text:
        return text

    # Tamil consonant range: U+0B95 to U+0BB9
    c_tamil = r'([\u0B95-\u0BB9])'
    # 1. ெ + Consonant + ா -> Consonant + ொ (U+0BCA)
    text = re.sub(r'\u0BC6' + c_tamil + r'\u0BBE', '\\g<1>\u0BCA', text)
    # 2. ே + Consonant + ா -> Consonant + ோ (U+0BCB)
    text = re.sub(r'\u0BC7' + c_tamil + r'\u0BBE', '\\g<1>\u0BCB', text)
    # 3. ெ + Consonant + ௗ -> Consonant + ௌ (U+0BCC)
    text = re.sub(r'\u0BC6' + c_tamil + r'\u0BD7', '\\g<1>\u0BCC', text)
    # 4. ெ + Consonant -> Consonant + ெ (U+0BC6)
    text = re.sub(r'\u0BC6' + c_tamil, '\\g<1>\u0BC6', text)
    # 5. ே + Consonant -> Consonant + ே (U+0BC7)
    text = re.sub(r'\u0BC7' + c_tamil, '\\g<1>\u0BC7', text)
    # 6. ை + Consonant -> Consonant + ை (U+0BC8)
    text = re.sub(r'\u0BC8' + c_tamil, '\\g<1>\u0BC8', text)

    # Malayalam pre-base vowels (U+0D46, U+0D47, U+0D48)
    c_mal = r'([\u0D15-\u0D3A])'
    text = re.sub(r'\u0D46' + c_mal + r'\u0D3E', '\\g<1>\u0D4A', text)
    text = re.sub(r'\u0D47' + c_mal + r'\u0D3E', '\\g<1>\u0D4B', text)
    text = re.sub(r'\u0D46' + c_mal, '\\g<1>\u0D46', text)
    text = re.sub(r'\u0D47' + c_mal, '\\g<1>\u0D47', text)
    text = re.sub(r'\u0D48' + c_mal, '\\g<1>\u0D48', text)

    # Bengali pre-base vowels (U+09C7, U+09C8)
    c_ben = r'([\u0995-\u09B9])'
    text = re.sub(r'\u09C7' + c_ben + r'\u09BE', '\\g<1>\u09CB', text)
    text = re.sub(r'\u09C7' + c_ben + r'\u09D7', '\\g<1>\u09CC', text)
    text = re.sub(r'\u09C7' + c_ben, '\\g<1>\u09C7', text)
    text = re.sub(r'\u09C8' + c_ben, '\\g<1>\u09C8', text)

    return text


def ensure_table_grids_in_docx(docx_path: str) -> bool:
    """
    Validates and repairs table grid structures (<w:tblGrid>) in word/document.xml.
    Prevents python-docx InvalidXmlError crashes and Word rendering bugs caused by
    missing table grids.
    """
    try:
        with zipfile.ZipFile(docx_path, 'r') as z:
            doc_xml = z.read('word/document.xml')
            if b'<w:tbl' not in doc_xml:
                return False

        root = ET.fromstring(doc_xml)
        w_ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        ns = {'w': w_ns}
        tables = root.findall('.//w:tbl', ns)
        modified = False

        for tbl in tables:
            grid = tbl.find('./w:tblGrid', ns)
            if grid is None:
                rows = tbl.findall('./w:tr', ns)
                max_cols = 1
                for r in rows:
                    cells = r.findall('./w:tc', ns)
                    if len(cells) > max_cols:
                        max_cols = len(cells)

                new_grid = ET.Element('{' + w_ns + '}tblGrid')
                col_width = max(720, int(9000 / max(1, max_cols)))
                for _ in range(max_cols):
                    gcol = ET.SubElement(new_grid, '{' + w_ns + '}gridCol')
                    gcol.set('{' + w_ns + '}w', str(col_width))

                tblPr = tbl.find('./w:tblPr', ns)
                if tblPr is not None:
                    idx = list(tbl).index(tblPr) + 1
                    tbl.insert(idx, new_grid)
                else:
                    tbl.insert(0, new_grid)
                modified = True

        if modified:
            tmp_dir = tempfile.mkdtemp()
            with zipfile.ZipFile(docx_path, 'r') as zin:
                zin.extractall(tmp_dir)

            doc_xml_path = os.path.join(tmp_dir, 'word', 'document.xml')
            ET.register_namespace('w', w_ns)
            root_bytes = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            with open(doc_xml_path, 'wb') as f_out:
                f_out.write(root_bytes)

            with zipfile.ZipFile(docx_path, 'w', zipfile.ZIP_DEFLATED) as zout:
                for foldername, _, filenames in os.walk(tmp_dir):
                    for filename in filenames:
                        filepath = os.path.join(foldername, filename)
                        arcname = os.path.relpath(filepath, tmp_dir)
                        zout.write(filepath, arcname)
            shutil.rmtree(tmp_dir)
            return True
    except Exception as exc:
        logger.warning(f"Notice during table grid validation: {exc}")
    return False


def clean_and_enhance_docx(docx_path: str) -> None:
    """
    Comprehensive DOCX post-processor:
    1. Guarantees valid table grids (<w:tblGrid>) in all tables.
    2. Replaces broken bullet artifacts (\\ufffd) with standard bullets.
    3. Reorders and normalizes Indic / Tamil character sequences into standard Unicode.
    4. Configures complex-script font metadata (w:cs="Nirmala UI") and language tags (ta-IN)
       while strictly preserving LTR direction for Indic scripts.
    5. Configures RTL paragraph and run properties for Arabic and Urdu.
    """
    if not os.path.exists(docx_path):
        return

    # 1. Guarantee valid table grids
    ensure_table_grids_in_docx(docx_path)

    try:
        doc = Document(docx_path)
        modified = False

        # Helper to process runs in a paragraph
        def process_paragraph_runs(p, is_table_cell=False):
            nonlocal modified
            p_text = p.text
            has_arabic = any('\u0600' <= c <= '\u06FF' or '\u0750' <= c <= '\u077F' for c in p_text)
            has_indic = any('\u0900' <= c <= '\u0D7F' for c in p_text)

            if has_arabic:
                # Set RTL paragraph properties
                pPr = p._p.get_or_add_pPr()
                bidi = pPr.find(qn('w:bidi'))
                if bidi is None:
                    bidi = OxmlElement('w:bidi')
                    pPr.append(bidi)
                modified = True

            # Run-boundary pre-base vowel transfer (e.g. vowel at end of run i, consonant at start of run i+1)
            pre_vowels = ('\u0BC6', '\u0BC7', '\u0BC8')
            c_tamil = r'^([\u0B95-\u0BB9])'
            runs = p.runs
            for idx in range(len(runs) - 1):
                r_curr = runs[idx]
                r_next = runs[idx + 1]
                for v in pre_vowels:
                    if r_curr.text.endswith(v):
                        m = re.match(c_tamil, r_next.text)
                        if m:
                            r_curr.text = r_curr.text[:-1]
                            r_next.text = v + r_next.text
                            modified = True
                            break

            for run in p.runs:
                r_text = run.text
                if not r_text:
                    continue

                # Fix bullet glyph artifacts
                if '\ufffd' in r_text:
                    r_text = r_text.replace('\ufffd', '• ')
                    run.text = r_text
                    modified = True

                # Normalize Indic visual ordering
                if any('\u0900' <= c <= '\u0D7F' for c in r_text):
                    norm = normalize_indic_text(r_text)
                    if norm != r_text:
                        run.text = norm
                        r_text = norm
                        modified = True

                    # Configure complex script font (Nirmala UI / Tamil)
                    rPr = run._r.get_or_add_rPr()
                    rFonts = rPr.find(qn('w:rFonts'))
                    if rFonts is None:
                        rFonts = OxmlElement('w:rFonts')
                        rPr.append(rFonts)
                    rFonts.set(qn('w:cs'), 'Nirmala UI')

                    # Set Indic language tag (preserve LTR)
                    lang = rPr.find(qn('w:lang'))
                    if lang is None:
                        lang = OxmlElement('w:lang')
                        rPr.append(lang)
                    # Use ta-IN if Tamil present
                    if any('\u0B80' <= c <= '\u0BFF' for c in r_text):
                        lang.set(qn('w:bidi'), 'ta-IN')
                    modified = True

                elif any('\u0600' <= c <= '\u06FF' for c in r_text):
                    # Configure Arabic/Urdu RTL run properties
                    rPr = run._r.get_or_add_rPr()
                    rtl = rPr.find(qn('w:rtl'))
                    if rtl is None:
                        rtl = OxmlElement('w:rtl')
                        rPr.append(rtl)
                    rFonts = rPr.find(qn('w:rFonts'))
                    if rFonts is None:
                        rFonts = OxmlElement('w:rFonts')
                        rPr.append(rFonts)
                    rFonts.set(qn('w:cs'), 'Arial')
                    lang = rPr.find(qn('w:lang'))
                    if lang is None:
                        lang = OxmlElement('w:lang')
                        rPr.append(lang)
                    lang.set(qn('w:bidi'), 'ar-SA')
                    modified = True

        for paragraph in doc.paragraphs:
            process_paragraph_runs(paragraph)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        process_paragraph_runs(paragraph, is_table_cell=True)

        if modified:
            doc.save(docx_path)
            logger.info(f"Enhanced typography and Indic metadata in: {docx_path}")
    except Exception as exc:
        logger.warning(f"DOCX enhancement notice: {exc}")


def clean_docx_artifacts(docx_path: str) -> None:
    """Legacy alias preserved for backward compatibility."""
    clean_and_enhance_docx(docx_path)


def inspect_pdf_document(pdf_path: str) -> Dict[str, Any]:
    """
    Inspects PDF document structure, encryption status, and text layer presence.
    """
    path_obj = Path(pdf_path)
    if not path_obj.exists() or path_obj.stat().st_size == 0:
        raise InvalidPdfError("The PDF document does not exist or is empty.")

    with open(pdf_path, 'rb') as f:
        header = f.read(5)
    if not header.startswith(b'%PDF-'):
        raise InvalidPdfError("The uploaded file does not have a valid PDF header.")

    try:
        doc = pymupdf.open(pdf_path)
    except Exception as err:
        import gc
        gc.collect()
        raise InvalidPdfError(f"Unable to parse PDF document: {str(err)}")

    if doc.is_encrypted and doc.needs_pass:
        doc.close()
        raise PasswordProtectedPdfError("This PDF is password-protected. Please remove the password before converting.")

    page_count = len(doc)
    if page_count == 0:
        doc.close()
        raise InvalidPdfError("The PDF document contains 0 pages.")

    is_scanned = True
    total_text_chars = 0
    sample_pages = min(page_count, 5)
    for i in range(sample_pages):
        sample_chars = len(doc[i].get_text().strip())
        total_text_chars += sample_chars
        if sample_chars > 0:
            is_scanned = False
            break

    doc.close()

    return {
        "page_count": page_count,
        "is_scanned": is_scanned,
        "text_chars": total_text_chars,
    }


def convert_pdf_to_docx(
    input_pdf_path: str,
    output_docx_path: Optional[str] = None
) -> Dict[str, Any]:
    """
    Converts a PDF document to Microsoft Word (.docx) using rule-based layout reconstruction.
    Preserves words, column boundaries, text styling, and page flow with optimized raster sampling.
    Applies automatic font CMap repair for corrupted Type0 fonts (e.g. Tamil / Indic) and
    enhances the resulting DOCX with valid table grids and complex-script font metadata.
    """
    start_time = time.perf_counter()

    # 1. Inspect and validate PDF
    meta = inspect_pdf_document(input_pdf_path)
    page_count = meta["page_count"]
    is_scanned = meta["is_scanned"]

    # 2. Determine target output path in isolated storage
    if not output_docx_path:
        filename = f"{uuid.uuid4().hex}_converted.docx"
        output_dir = getattr(settings, 'OUTPUTS_DIR', Path(input_pdf_path).parent)
        output_dir.mkdir(parents=True, exist_ok=True)
        output_docx_path = str(output_dir / filename)

    # 3. Check and repair corrupted embedded font CMaps (e.g. Tamil Identity-H subsets)
    repaired_pdf_path, is_repaired = repair_pdf_cmaps(str(input_pdf_path))
    pdf_to_convert = repaired_pdf_path if is_repaired else str(input_pdf_path)

    # 4. Perform conversion via pdf2docx Converter with optimized sampling
    cv = None
    try:
        cv = Converter(str(pdf_to_convert))
        cv.default_settings['clip_image_res_ratio'] = 1.25
        cv.default_settings['float_image_ignorable_gap'] = 10.0
        cv.default_settings['min_svg_gap_dx'] = 25.0
        cv.default_settings['ignore_page_error'] = True
        cv.default_settings['debug'] = False
        cv.convert(str(output_docx_path))
    except Exception as err:
        logger.error(f"pdf2docx conversion error: {err}", exc_info=True)
        raise WordEngineException(f"Failed to reconstruct DOCX layout: {str(err)}")
    finally:
        if cv is not None:
            try:
                cv.close()
            except Exception:
                pass
        # Clean up temporary repaired PDF if one was created
        if is_repaired and os.path.exists(repaired_pdf_path) and repaired_pdf_path != str(input_pdf_path):
            try:
                os.remove(repaired_pdf_path)
            except Exception:
                pass

    # 5. Clean artifacts, guarantee valid table grids, and enhance typography metadata
    clean_and_enhance_docx(str(output_docx_path))

    elapsed_ms = int((time.perf_counter() - start_time) * 1000)
    file_size = os.path.getsize(output_docx_path)

    return {
        "success": True,
        "output_path": str(output_docx_path),
        "output_filename": Path(output_docx_path).name,
        "page_count": page_count,
        "size_bytes": file_size,
        "is_scanned": is_scanned,
        "execution_time_ms": elapsed_ms,
    }
